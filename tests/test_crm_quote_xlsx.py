import sys
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from routes import crm


class QuoteXlsxRowHeightTest(unittest.TestCase):
    def setUp(self):
        self.template_path = ROOT / "data" / "立翔112年上半年工程.xlsx"

    @staticmethod
    def _quote_with_note(note: str):
        item = SimpleNamespace(
            id=1,
            sort_order=0,
            description="配電盤控制線施工",
            unit="式",
            quantity=1,
            unit_price=1000,
            amount=1000,
            note=note,
        )
        return SimpleNamespace(
            recipient_name="嘉義市四維南路Ａ停車場",
            site_address="嘉義市四維南路",
            issue_date=date(2026, 7, 15),
            items=[item],
            total_amount=1000,
            subtotal=1000,
        )

    def test_long_cjk_note_wraps_and_increases_item_row_height(self):
        workbook = load_workbook(self.template_path)
        self.addCleanup(workbook.close)
        worksheet = workbook.worksheets[0]
        original_height = worksheet.row_dimensions[7].height

        crm._apply_quote_to_template_sheet(
            worksheet,
            self._quote_with_note("含盤內控制線等白螺絲配電工膠帶等"),
            None,
            None,
        )

        self.assertTrue(worksheet["J7"].alignment.wrap_text)
        self.assertGreater(worksheet.row_dimensions[7].height, original_height)

    def test_short_note_keeps_template_minimum_row_height(self):
        workbook = load_workbook(self.template_path)
        self.addCleanup(workbook.close)
        worksheet = workbook.worksheets[0]
        original_height = worksheet.row_dimensions[7].height

        crm._apply_quote_to_template_sheet(
            worksheet,
            self._quote_with_note("5*2C"),
            None,
            None,
        )

        self.assertEqual(worksheet.row_dimensions[7].height, original_height)

    def test_long_rows_use_safe_vertical_paging_and_keep_total_row(self):
        workbook = load_workbook(self.template_path)
        self.addCleanup(workbook.close)
        worksheet = workbook.worksheets[0]

        crm._apply_quote_to_template_sheet(
            worksheet,
            self._quote_with_note("含盤內控制線等白螺絲配電工膠帶等"),
            None,
            None,
        )

        self.assertTrue(worksheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(worksheet.page_setup.fitToWidth, 1)
        self.assertEqual(worksheet.page_setup.fitToHeight, 0)
        self.assertIsNone(worksheet.page_setup.scale)
        self.assertEqual(worksheet.print_title_rows, "$1:$6")
        self.assertEqual(worksheet["C27"].value, "總計")
        self.assertEqual(worksheet["I27"].value, 1000)


if __name__ == "__main__":
    unittest.main()
