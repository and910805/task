"""Cross-workspace (tenant) isolation tests against the full application."""

import io
import os
import sys
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

PNG_BYTES = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
    b"\x00\x00\x00\rIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82"
)
PASSWORD = "correct-horse-battery"


def _build_app(directory: str):
    env = {
        "DATABASE_URL": f"sqlite:///{Path(directory) / 'test.db'}",
        "SECRET_KEY": "test-secret",
        "JWT_SECRET_KEY": "test-jwt-secret",
        "UPLOAD_FOLDER": str(Path(directory) / "uploads"),
        "INIT_DB_ON_STARTUP": "1",
        "ALLOW_PUBLIC_SIGNUP": "1",
    }
    with patch.dict(os.environ, env, clear=False):
        from app import create_app

        return create_app()


def _clear_rate_limits():
    from rate_limit import _BUCKETS, _LOCK

    with _LOCK:
        _BUCKETS.clear()


class WorkspaceTestCase(unittest.TestCase):
    def setUp(self):
        _clear_rate_limits()
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.app = _build_app(self.temp_dir.name)
        self.client = self.app.test_client()

    def tearDown(self):
        from extensions import db

        with self.app.app_context():
            db.session.remove()
            db.engine.dispose()
        self.temp_dir.cleanup()

    # -- helpers ---------------------------------------------------------
    def signup(self, username, company, industry="plumbing_electrical"):
        response = self.client.post(
            "/api/auth/signup",
            json={
                "username": username,
                "password": PASSWORD,
                "company_name": company,
                "industry": industry,
                "accept_terms": True,
            },
        )
        self.assertEqual(response.status_code, 201, response.get_json())
        data = response.get_json()
        return {"token": data["token"], "workspace_id": data["active_workspace_id"], "user": data["user"]}

    def headers(self, session, workspace_id=None):
        headers = {"Authorization": f"Bearer {session['token']}"}
        ws = workspace_id if workspace_id is not None else session.get("workspace_id")
        if ws is not None:
            headers["X-Workspace-Id"] = str(ws)
        return headers

    def invite(self, owner, role="worker", **extra):
        response = self.client.post(
            "/api/workspaces/current/invitations",
            json={"role": role, **extra},
            headers=self.headers(owner),
        )
        self.assertEqual(response.status_code, 201, response.get_json())
        return response.get_json()["code"]

    def join_new_user(self, owner, username, role="worker"):
        code = self.invite(owner, role=role)
        response = self.client.post(
            "/api/auth/register",
            json={"username": username, "password": PASSWORD, "invite_code": code},
        )
        self.assertEqual(response.status_code, 201, response.get_json())
        data = response.get_json()
        self.assertEqual(data["active_workspace_id"], owner["workspace_id"])
        self.assertEqual(data["user"]["role"], role)
        return {"token": data["token"], "workspace_id": data["active_workspace_id"], "user": data["user"]}

    def create_task(self, session, assignee_ids=(), title="更換熱水器"):
        response = self.client.post(
            "/api/tasks/create",
            json={
                "title": title,
                "description": "現場更換",
                "location": "台北市",
                "expected_time": (datetime.utcnow() + timedelta(hours=2)).isoformat(),
                "status": "尚未接單",
                "assignee_ids": list(assignee_ids),
            },
            headers=self.headers(session),
        )
        self.assertEqual(response.status_code, 201, response.get_json())
        return response.get_json()


class CrossWorkspaceIsolationTest(WorkspaceTestCase):
    def setUp(self):
        super().setUp()
        self.owner_a = self.signup("owner-a", "甲水電行")
        self.owner_b = self.signup("owner-b", "乙清潔", industry="cleaning")
        self.worker_a = self.join_new_user(self.owner_a, "worker-a")
        self.worker_b = self.join_new_user(self.owner_b, "worker-b")
        self.task_a = self.create_task(self.owner_a, [self.worker_a["user"]["id"]])

    def test_task_list_only_contains_own_workspace(self):
        own = self.client.get("/api/tasks/", headers=self.headers(self.owner_a)).get_json()
        other = self.client.get("/api/tasks/", headers=self.headers(self.owner_b)).get_json()
        self.assertEqual([task["id"] for task in own], [self.task_a["id"]])
        self.assertEqual(other, [])
        available = self.client.get("/api/tasks/?available=1", headers=self.headers(self.worker_b)).get_json()
        self.assertEqual(available, [])

    def test_other_workspace_task_is_not_found_for_every_task_endpoint(self):
        task_id = self.task_a["id"]
        intruders = (self.owner_b, self.worker_b)
        calls = [
            ("get", f"/api/tasks/{task_id}", None),
            ("put", f"/api/tasks/{task_id}", {"title": "hijack"}),
            ("patch", f"/api/tasks/update/{task_id}", {"status": "進行中"}),
            ("post", f"/api/tasks/{task_id}/updates", {"note": "hi"}),
            ("post", f"/api/tasks/{task_id}/time/start", {}),
            ("post", f"/api/tasks/{task_id}/time/stop", {}),
            ("post", f"/api/tasks/{task_id}/time/manual", {"user_ids": [1], "work_hours": 1}),
            ("post", f"/api/tasks/{task_id}/accept", {}),
            ("post", f"/api/tasks/{task_id}/assignees/add", {"assignee_ids": [self.worker_b["user"]["id"]]}),
            ("post", f"/api/tasks/{task_id}/archive", {}),
            ("post", f"/api/tasks/{task_id}/restore", {}),
            ("delete", f"/api/tasks/{task_id}", None),
            ("post", f"/api/upload/tasks/{task_id}/signature", {"data_url": "data:image/png;base64,AAAA"}),
        ]
        for intruder in intruders:
            for method, url, body in calls:
                with self.subTest(user=intruder["user"]["username"], method=method, url=url):
                    response = getattr(self.client, method)(url, json=body, headers=self.headers(intruder))
                    self.assertIn(response.status_code, {403, 404}, response.get_json())
        photo = self.client.post(
            f"/api/upload/tasks/{task_id}/images",
            data={"file": (io.BytesIO(PNG_BYTES), "a.png")},
            headers=self.headers(self.owner_b),
            content_type="multipart/form-data",
        )
        self.assertEqual(photo.status_code, 404)
        detail = self.client.get(f"/api/tasks/{task_id}", headers=self.headers(self.owner_a)).get_json()
        self.assertEqual(detail["title"], "更換熱水器")
        self.assertEqual(detail["updates"], [])

    def test_workspace_header_for_foreign_workspace_is_forbidden(self):
        response = self.client.get(
            "/api/tasks/", headers=self.headers(self.owner_b, workspace_id=self.owner_a["workspace_id"])
        )
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()["code"], "workspace_forbidden")

    def test_cannot_assign_or_list_members_of_another_workspace(self):
        response = self.client.post(
            "/api/tasks/create",
            json={
                "title": "x",
                "description": "x",
                "location": "x",
                "expected_time": datetime.utcnow().isoformat(),
                "status": "尚未接單",
                "assignee_ids": [self.worker_b["user"]["id"]],
            },
            headers=self.headers(self.owner_a),
        )
        self.assertEqual(response.status_code, 404)
        members = self.client.get("/api/auth/users", headers=self.headers(self.owner_a)).get_json()["users"]
        self.assertEqual({m["username"] for m in members}, {"owner-a", "worker-a"})
        assignable = self.client.get("/api/auth/assignable-users", headers=self.headers(self.owner_a)).get_json()
        self.assertNotIn("worker-b", {m["username"] for m in assignable})
        for method, url, body in (
            ("put", f"/api/auth/users/{self.worker_b['user']['id']}", {"role": "admin"}),
            ("put", f"/api/auth/users/{self.worker_b['user']['id']}", {"password": "new-password-123"}),
            ("delete", f"/api/auth/users/{self.worker_b['user']['id']}", None),
        ):
            with self.subTest(method=method, body=body):
                response = getattr(self.client, method)(url, json=body, headers=self.headers(self.owner_a))
                self.assertEqual(response.status_code, 404)

    def test_schedule_conflicts_do_not_leak_other_workspace_tasks(self):
        # Same wall-clock slot in another company must not collide or leak titles.
        response = self.client.post(
            "/api/tasks/create",
            json={
                "title": "乙公司案件",
                "description": "x",
                "location": "x",
                "expected_time": self.task_a["expected_time"],
                "status": "尚未接單",
                "assignee_ids": [self.worker_b["user"]["id"]],
            },
            headers=self.headers(self.owner_b),
        )
        self.assertEqual(response.status_code, 201, response.get_json())

    def test_attachments_and_reports_are_not_downloadable_across_workspaces(self):
        upload = self.client.post(
            f"/api/upload/tasks/{self.task_a['id']}/images",
            data={"file": (io.BytesIO(PNG_BYTES), "site.png")},
            headers=self.headers(self.worker_a),
            content_type="multipart/form-data",
        )
        self.assertEqual(upload.status_code, 201, upload.get_json())
        url = upload.get_json()["url"]
        path = url.split("?")[0]
        own = self.client.get(path, headers=self.headers(self.owner_a))
        self.assertEqual(own.status_code, 200)
        other = self.client.get(path, headers=self.headers(self.owner_b))
        self.assertEqual(other.status_code, 404)

        report_b = self.client.get("/api/export/tasks", headers=self.headers(self.owner_b)).get_json()
        report_a = self.client.get("/api/export/tasks", headers=self.headers(self.owner_a)).get_json()
        self.assertIn(f"ws{self.owner_b['workspace_id']}", report_b["url"])
        path_a = report_a["url"].split("?")[0]
        self.assertEqual(self.client.get(path_a, headers=self.headers(self.owner_a)).status_code, 200)
        self.assertEqual(self.client.get(path_a, headers=self.headers(self.owner_b)).status_code, 404)

        from openpyxl import load_workbook

        body = self.client.get(report_b["url"].split("?")[0], headers=self.headers(self.owner_b)).data
        sheet = load_workbook(io.BytesIO(body))["Tasks"]
        titles = [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)]
        self.assertNotIn("更換熱水器", titles)

    def test_site_locations_and_settings_are_per_workspace(self):
        for owner in (self.owner_a, self.owner_b):
            response = self.client.post("/api/site-locations/", json={"name": "公司倉庫"}, headers=self.headers(owner))
            self.assertEqual(response.status_code, 201, response.get_json())
        locations_b = self.client.get("/api/site-locations/", headers=self.headers(self.owner_b)).get_json()
        self.assertEqual(len(locations_b), 1)
        location_a = self.client.get("/api/site-locations/", headers=self.headers(self.owner_a)).get_json()[0]
        self.assertEqual(
            self.client.delete(f"/api/site-locations/{location_a['id']}", headers=self.headers(self.owner_b)).status_code,
            404,
        )

        self.client.put("/api/settings/roles/worker", json={"label": "清潔夥伴"}, headers=self.headers(self.owner_b))
        labels_a = self.client.get("/api/settings/roles", headers=self.headers(self.owner_a)).get_json()["labels"]
        labels_b = self.client.get("/api/settings/roles", headers=self.headers(self.owner_b)).get_json()["labels"]
        self.assertEqual(labels_a["worker"], "水電師傅")
        self.assertEqual(labels_b["worker"], "清潔夥伴")

        self.client.put("/api/settings/branding/name", json={"name": "乙清潔有限公司"}, headers=self.headers(self.owner_b))
        self.assertEqual(
            self.client.get("/api/settings/branding", headers=self.headers(self.owner_a)).get_json()["name"], "甲水電行"
        )
        self.assertEqual(self.client.get("/api/settings/branding").get_json()["name"], "TaskGo")

        self.client.put("/api/settings/task-update-templates", json={"templates": ["乙專用"]}, headers=self.headers(self.owner_b))
        templates_a = self.client.get("/api/settings/task-update-templates", headers=self.headers(self.owner_a)).get_json()
        self.assertNotIn("乙專用", templates_a["templates"])

    def test_unscoped_modules_are_closed_for_new_workspaces(self):
        for url in ("/api/settings/notifications/email",):
            with self.subTest(url=url):
                response = self.client.get(url, headers=self.headers(self.owner_a))
                self.assertEqual(response.status_code, 403)
                self.assertEqual(response.get_json()["code"], "module_unavailable")

    def test_scoped_business_modules_are_available_for_new_workspaces(self):
        for url in ("/api/crm/customers", "/api/crm/boot", "/api/materials/items"):
            with self.subTest(url=url):
                response = self.client.get(url, headers=self.headers(self.owner_a))
                self.assertEqual(response.status_code, 200, response.get_json())

    def test_role_and_membership_changes_apply_immediately(self):
        worker_headers = self.headers(self.worker_a)
        self.assertEqual(self.client.post("/api/tasks/create", json={}, headers=worker_headers).status_code, 403)
        promote = self.client.put(
            f"/api/auth/users/{self.worker_a['user']['id']}", json={"role": "site_supervisor"}, headers=self.headers(self.owner_a)
        )
        self.assertEqual(promote.status_code, 200)
        self.assertEqual(self.create_task(self.worker_a, title="主管建立")["title"], "主管建立")
        remove = self.client.delete(f"/api/auth/users/{self.worker_a['user']['id']}", headers=self.headers(self.owner_a))
        self.assertEqual(remove.status_code, 200)
        self.assertEqual(self.client.get("/api/tasks/", headers=worker_headers).status_code, 401)

    def test_admin_cannot_escalate_or_touch_owner(self):
        admin_code = self.invite(self.owner_a, role="admin")
        admin = self.client.post(
            "/api/auth/register", json={"username": "admin-a", "password": PASSWORD, "invite_code": admin_code}
        ).get_json()
        admin_session = {"token": admin["token"], "workspace_id": self.owner_a["workspace_id"]}
        owner_id = self.owner_a["user"]["id"]
        self.assertEqual(
            self.client.post("/api/workspaces/current/invitations", json={"role": "admin"}, headers=self.headers(admin_session)).status_code,
            403,
        )
        self.assertEqual(
            self.client.put(f"/api/auth/users/{owner_id}", json={"role": "worker"}, headers=self.headers(admin_session)).status_code,
            400,
        )
        self.assertEqual(
            self.client.delete(f"/api/auth/users/{owner_id}", headers=self.headers(admin_session)).status_code, 400
        )
        self.assertEqual(
            self.client.post("/api/workspaces/current/transfer-ownership", json={"user_id": admin["user"]["id"]}, headers=self.headers(admin_session)).status_code,
            403,
        )

    def test_invitations_are_scoped_and_single_use(self):
        code = self.invite(self.owner_a)
        invitations_b = self.client.get("/api/workspaces/current/invitations", headers=self.headers(self.owner_b)).get_json()
        self.assertTrue(all(item["code_hint"] != code[-4:] for item in invitations_b["invitations"]) or invitations_b["invitations"] == [])
        first = self.client.post("/api/auth/register", json={"username": "u1", "password": PASSWORD, "invite_code": code})
        self.assertEqual(first.status_code, 201)
        second = self.client.post("/api/auth/register", json={"username": "u2", "password": PASSWORD, "invite_code": code})
        self.assertEqual(second.status_code, 404)
        anonymous = self.client.post("/api/auth/register", json={"username": "u3", "password": PASSWORD})
        self.assertEqual(anonymous.status_code, 403)
        invitation_id = self.client.get(
            "/api/workspaces/current/invitations", headers=self.headers(self.owner_a)
        ).get_json()["invitations"][0]["id"]
        self.assertEqual(
            self.client.delete(f"/api/workspaces/current/invitations/{invitation_id}", headers=self.headers(self.owner_b)).status_code,
            404,
        )

    def test_user_in_two_workspaces_switches_explicitly(self):
        code = self.invite(self.owner_b)
        joined = self.client.post("/api/workspaces/join", json={"code": code}, headers=self.headers(self.worker_a))
        self.assertEqual(joined.status_code, 200, joined.get_json())
        self.assertEqual(len(joined.get_json()["workspaces"]), 2)
        in_a = self.client.get("/api/tasks/", headers=self.headers(self.worker_a, self.owner_a["workspace_id"])).get_json()
        in_b = self.client.get("/api/tasks/", headers=self.headers(self.worker_a, self.owner_b["workspace_id"])).get_json()
        self.assertEqual([t["id"] for t in in_a], [self.task_a["id"]])
        self.assertEqual(in_b, [])

    def test_account_deletion_requires_ownership_transfer(self):
        response = self.client.delete("/api/auth/account", json={"password": PASSWORD}, headers=self.headers(self.owner_a))
        self.assertEqual(response.status_code, 409)
        lone = self.signup("solo", "一人公司")
        self.create_task(lone, title="solo task")
        deleted = self.client.delete("/api/auth/account", json={"password": PASSWORD}, headers=self.headers(lone))
        self.assertEqual(deleted.status_code, 200)
        self.assertEqual(
            self.client.post("/api/auth/login", json={"username": "solo", "password": PASSWORD}).status_code, 401
        )


class DispatchFlowTest(WorkspaceTestCase):
    def test_manager_dispatches_and_worker_completes_with_photo(self):
        owner = self.signup("boss", "丙維修", industry="repair")
        worker = self.join_new_user(owner, "tech")
        task = self.create_task(owner, [worker["user"]["id"]])
        worker_headers = self.headers(worker)

        mine = self.client.get("/api/tasks/", headers=worker_headers).get_json()
        self.assertEqual([t["id"] for t in mine], [task["id"]])
        self.assertEqual(mine[0]["assignees"][0]["role_label"], "維修技師")

        self.assertEqual(self.client.post(f"/api/tasks/{task['id']}/updates", json={"status": "進行中"}, headers=worker_headers).status_code, 201)
        self.assertEqual(self.client.post(f"/api/tasks/{task['id']}/time/start", headers=worker_headers).status_code, 201)
        blocked = self.client.post(f"/api/tasks/{task['id']}/updates", json={"status": "已完成", "note": "完成"}, headers=worker_headers)
        self.assertEqual(blocked.status_code, 400)
        photo = self.client.post(
            f"/api/upload/tasks/{task['id']}/images",
            data={"file": (io.BytesIO(PNG_BYTES), "done.png"), "note": "完工照"},
            headers=worker_headers,
            content_type="multipart/form-data",
        )
        self.assertEqual(photo.status_code, 201)
        self.assertEqual(self.client.post(f"/api/tasks/{task['id']}/time/stop", headers=worker_headers).status_code, 200)
        done = self.client.post(f"/api/tasks/{task['id']}/updates", json={"status": "已完成", "note": "已修復並測試"}, headers=worker_headers)
        self.assertEqual(done.status_code, 201, done.get_json())

        detail = self.client.get(f"/api/tasks/{task['id']}", headers=self.headers(owner)).get_json()
        self.assertEqual(detail["status"], "已完成")
        self.assertIsNotNone(detail["completed_at"])
        self.assertEqual(len([a for a in detail["attachments"] if a["file_type"] == "image"]), 1)

    def test_session_payload_reports_workspace_role(self):
        owner = self.signup("boss2", "丁裝修", industry="renovation")
        me = self.client.get("/api/auth/me", headers=self.headers(owner)).get_json()
        self.assertEqual(me["role"], "admin")
        self.assertTrue(me["is_owner"])
        self.assertEqual(me["workspaces"][0]["name"], "丁裝修")
        for module in ("crm", "materials", "reports"):
            self.assertTrue(me["workspaces"][0]["modules"][module])
        self.assertFalse(me["workspaces"][0]["modules"]["line_settings"])


class PushNotificationTest(WorkspaceTestCase):
    def test_device_registration_follows_the_signed_in_user(self):
        from models import DeviceToken

        first = self.signup("push-a", "推播甲")
        second = self.signup("push-b", "推播乙")
        self.assertEqual(self.client.get("/api/workspaces/push-status", headers=self.headers(first)).get_json(), {"configured": False})
        for session in (first, second):
            response = self.client.post("/api/workspaces/devices", json={"token": "abc123", "platform": "ios"}, headers=self.headers(session))
            self.assertEqual(response.status_code, 200)
        with self.app.app_context():
            rows = DeviceToken.query.all()
            self.assertEqual([(row.token, row.user_id) for row in rows], [("abc123", second["user"]["id"])])
        self.client.delete("/api/workspaces/devices", json={"token": "abc123"}, headers=self.headers(first))
        with self.app.app_context():
            self.assertEqual(DeviceToken.query.count(), 1)
        self.client.delete("/api/workspaces/devices", json={"token": "abc123"}, headers=self.headers(second))
        with self.app.app_context():
            self.assertEqual(DeviceToken.query.count(), 0)

    def test_apns_payload_and_invalid_token_detection(self):
        from cryptography.hazmat.primitives import serialization
        from cryptography.hazmat.primitives.asymmetric import ec

        from services import push

        key = ec.generate_private_key(ec.SECP256R1()).private_bytes(
            serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()
        ).decode()
        env = {"APNS_KEY_ID": "KEYID12345", "APNS_TEAM_ID": "TEAMID1234", "APNS_BUNDLE_ID": "online.kuanlin.taskgo", "APNS_PRIVATE_KEY": key}

        class FakeResponse:
            def __init__(self, status, reason=""):
                self.status_code = status
                self._reason = reason

            def json(self):
                return {"reason": self._reason}

        class FakeClient:
            def __init__(self):
                self.calls = []

            def post(self, url, json, headers):
                self.calls.append((url, json, headers))
                return FakeResponse(200) if url.endswith("/good") else FakeResponse(400, "BadDeviceToken")

        with patch.dict(os.environ, env, clear=False), self.app.app_context():
            push._token_cache.clear()
            self.assertTrue(push.apns_configured())
            payload = push.build_payload("新派工：換水龍頭", "台北市", {"task_id": 7, "workspace_id": 3})
            client = FakeClient()
            invalid = push.send_to_tokens(["good", "stale"], payload, client=client)
        self.assertEqual(invalid, ["stale"])
        url, body, headers = client.calls[0]
        self.assertEqual(url, "https://api.push.apple.com/3/device/good")
        self.assertEqual(body["task_id"], 7)
        self.assertEqual(body["aps"]["alert"]["title"], "新派工：換水龍頭")
        self.assertEqual(headers["apns-topic"], "online.kuanlin.taskgo")
        self.assertTrue(headers["authorization"].startswith("bearer "))


class AppStoreSupportTest(WorkspaceTestCase):
    def test_universal_links_file_only_when_team_id_configured(self):
        self.assertEqual(self.client.get("/.well-known/apple-app-site-association").status_code, 404)
        with patch.dict(os.environ, {"APPLE_TEAM_ID": "TEAMID1234"}, clear=False):
            response = self.client.get("/.well-known/apple-app-site-association")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["Content-Type"], "application/json")
        details = response.get_json()["applinks"]["details"][0]
        self.assertEqual(details["appID"], "TEAMID1234.online.kuanlin.taskgo")
        self.assertEqual([c["/"] for c in details["components"]], ["/tasks/*", "/join"])

    def test_native_origin_preflight_allows_workspace_header(self):
        response = self.client.options(
            "/api/tasks/",
            headers={
                "Origin": "capacitor://localhost",
                "Access-Control-Request-Method": "GET",
                "Access-Control-Request-Headers": "authorization,x-workspace-id",
            },
        )
        self.assertEqual(response.headers.get("Access-Control-Allow-Origin"), "capacitor://localhost")
        self.assertIn("x-workspace-id", response.headers.get("Access-Control-Allow-Headers", "").lower())

    def test_review_workspace_is_isolated_and_password_is_not_echoed(self):
        from click.testing import CliRunner

        runner = CliRunner()
        # Run inside this app's context so the command uses the test database.
        with patch.dict(os.environ, {"REVIEW_ACCOUNT_PASSWORD": "review-demo-pass-1"}, clear=False), self.app.app_context():
            result = runner.invoke(self.app.cli, ["create-review-workspace"], catch_exceptions=False)
        self.assertEqual(result.exit_code, 0, result.output)
        self.assertNotIn("review-demo-pass-1", result.output)
        login = self.client.post("/api/auth/login", json={"username": "review-owner", "password": "review-demo-pass-1"}).get_json()
        headers = {"Authorization": f"Bearer {login['token']}"}
        tasks = self.client.get("/api/tasks/", headers=headers).get_json()
        self.assertEqual(len(tasks), 4)
        self.assertEqual(login["workspaces"][0]["role_label"], "擁有者")
        worker = self.client.post("/api/auth/login", json={"username": "review-worker", "password": "review-demo-pass-1"}).get_json()
        self.assertEqual(worker["user"]["role_label"], "水電師傅")
        other = self.signup("outsider", "別家公司")
        self.assertEqual(self.client.get("/api/tasks/", headers=self.headers(other)).get_json(), [])
        with patch.dict(os.environ, {"REVIEW_ACCOUNT_PASSWORD": "review-demo-pass-1"}, clear=False), self.app.app_context():
            again = runner.invoke(self.app.cli, ["create-review-workspace"])
        self.assertNotEqual(again.exit_code, 0)


class RouteTenancyPolicyTest(WorkspaceTestCase):
    def test_every_api_route_declares_a_tenancy_policy(self):
        from tenancy import POLICY_ATTR

        missing = []
        for rule in self.app.url_map.iter_rules():
            if not rule.rule.startswith("/api/"):
                continue
            policy = getattr(self.app.view_functions[rule.endpoint], POLICY_ATTR, None)
            if policy not in {"public", "account", "workspace"}:
                missing.append(rule.rule)
        self.assertEqual(missing, [])

    def test_public_routes_are_an_explicit_allowlist(self):
        from tenancy import POLICY_ATTR

        public = sorted(
            rule.rule
            for rule in self.app.url_map.iter_rules()
            if rule.rule.startswith("/api/")
            and getattr(self.app.view_functions[rule.endpoint], POLICY_ATTR, None) == "public"
        )
        self.assertEqual(
            public,
            sorted(
                [
                    "/api/auth/login",
                    "/api/auth/logout",
                    "/api/auth/register",
                    "/api/auth/signup",
                    "/api/crm/public/bookings",
                    "/api/health",
                    "/api/line/webhook",
                    "/api/line/webhook/public",
                    "/api/public/photos",
                    "/api/settings/branding",
                    "/api/settings/branding/logo/<string:token>",
                    "/api/upload/files/<path:filename>",
                    "/api/workspaces/invitations/preview",
                ]
            ),
        )


class LegacyMigrationTest(unittest.TestCase):
    """Upgrade/verify/downgrade round trip on a single-tenant style database."""

    def setUp(self):
        _clear_rate_limits()
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.app = _build_app(self.temp_dir.name)

    def tearDown(self):
        from extensions import db

        with self.app.app_context():
            db.session.remove()
            db.engine.dispose()
        self.temp_dir.cleanup()

    def _seed_legacy_data(self):
        from extensions import db
        from models import RoleLabel, SiteLocation, SiteSetting, Task, User

        admin = User(username="lx-admin", role="admin")
        admin.set_password(PASSWORD)
        worker = User(username="lx-worker", role="worker")
        worker.set_password(PASSWORD)
        db.session.add_all([admin, worker])
        db.session.flush()
        db.session.add(
            Task(
                title="舊任務",
                description="d",
                location="嘉義",
                expected_time=datetime.utcnow(),
                assigned_to_id=worker.id,
                assigned_by_id=admin.id,
            )
        )
        db.session.add(SiteLocation(name="總部"))
        db.session.add(SiteSetting(key="branding_name", value="立翔水電工程行"))
        db.session.add(RoleLabel(role="worker", label="水電工"))
        db.session.commit()

    def test_round_trip(self):
        import workspace_migration as migration
        from extensions import db

        with self.app.app_context():
            # Start from the single-tenant shape: no workspace tables/data.
            migration.downgrade()
            self._seed_legacy_data()
            before = migration.inventory()["tables"]
            self.assertFalse(migration.verify()["ok"])

            report = migration.upgrade("立翔水電行")
            self.assertTrue(report["first_run"])
            self.assertEqual(report["memberships_created"], 2)
            self.assertEqual(report["backfilled"]["task"], 1)
            verified = migration.verify()
            self.assertTrue(verified["ok"], verified)
            again = migration.upgrade("立翔水電行")
            self.assertFalse(again["first_run"])
            self.assertEqual(again["memberships_created"], 0)
            after = migration.inventory()["tables"]
            for table, count in before.items():
                self.assertEqual(after[table], count, table)
            db.session.remove()

        client = self.app.test_client()
        login = client.post("/api/auth/login", json={"username": "lx-admin", "password": PASSWORD}).get_json()
        headers = {"Authorization": f"Bearer {login['token']}"}
        self.assertEqual(login["user"]["role"], "admin")
        self.assertTrue(login["workspaces"][0]["modules"]["crm"])
        self.assertEqual(client.get("/api/tasks/", headers=headers).get_json()[0]["title"], "舊任務")
        self.assertEqual(client.get("/api/settings/branding", headers=headers).get_json()["name"], "立翔水電工程行")
        labels = client.get("/api/settings/roles", headers=headers).get_json()["labels"]
        self.assertEqual(labels["worker"], "水電工")
        self.assertEqual(labels["site_supervisor"], "現場主管")
        self.assertEqual(client.get("/api/crm/customers", headers=headers).status_code, 200)

        other = client.post(
            "/api/auth/signup",
            json={"username": "new-co", "password": PASSWORD, "company_name": "新公司", "accept_terms": True},
        ).get_json()
        client.post(
            "/api/tasks/create",
            json={"title": "新公司任務", "description": "d", "location": "x", "expected_time": datetime.utcnow().isoformat(), "status": "尚未接單"},
            headers={"Authorization": f"Bearer {other['token']}"},
        )

        with self.app.app_context():
            with self.assertRaises(migration.MigrationError):
                migration.downgrade()
            result = migration.downgrade(force_discard_other_workspaces=True)
            self.assertEqual(result["accounts_outside_legacy"], 1)
            final = migration.inventory()
            self.assertFalse(final["migrated"])
            self.assertEqual(final["tables"]["task"], before["task"])
            self.assertEqual(final["tables"]["user"], before["user"])
            self.assertEqual(final["roles"], {"admin": 1, "worker": 1})


if __name__ == "__main__":
    unittest.main()
